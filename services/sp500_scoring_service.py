from __future__ import annotations

import json
import logging
import time
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from io import StringIO
from pathlib import Path
from typing import Any

import httpx
import pandas as pd

from core.config import get_settings, load_yaml_config
from core.models import PortfolioExposureModel, TechnicalSnapshotModel
from data.database import AssetORM
from market_regime.regime_service import MarketRegimeService
from services.rebalance_service import RebalanceService
from services.recommendation_service import RecommendationService
from services.risk_service import RiskService
from services.scoring_service import ScoringService
from services.support_detection_service import SupportDetectionService
from services.technical_service import TechnicalService

logger = logging.getLogger(__name__)


@dataclass
class SP500ScoringResult:
    ranking: pd.DataFrame
    errors: pd.DataFrame
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass
class SP500IncrementalRefreshResult:
    target_session: date
    symbols_total: int
    symbols_current: int
    symbols_refreshed: int
    symbols_failed: int
    errors: list[dict[str, str]] = field(default_factory=list)


class SP500ScoringService:
    """Scores an external S&P 500 snapshot without mutating the tracked universe."""

    def __init__(
        self,
        config: dict[str, Any] | None = None,
        *,
        now: datetime | None = None,
    ) -> None:
        self.config = config or load_yaml_config("sp500_scoring.yaml")
        self.now = now or datetime.now(UTC)
        self.root_dir = get_settings().root_dir
        self.technical = TechnicalService()
        self.support = SupportDetectionService()
        self.scoring = ScoringService()
        self.risk = RiskService()
        self.rebalance = RebalanceService()
        self.recommendation = RecommendationService()
        self.regime = MarketRegimeService()

    def refresh_price_cache_incremental(self) -> SP500IncrementalRefreshResult:
        """Append recent bars only for constituents whose cache is behind."""
        constituents = self.load_constituents(force=False)
        symbols = constituents["yahoo_symbol"].dropna().astype(str).tolist()
        target_session = self._expected_market_session()
        current = 0
        pending: list[str] = []
        bootstrap: list[str] = []
        existing_frames: dict[str, pd.DataFrame] = {}
        minimum_rows = int(
            self.config.get("market_data", {}).get("min_history_rows", 220)
        )
        for symbol in symbols:
            latest, cached_rows = self._cached_price_status(symbol)
            if latest is None:
                bootstrap.append(symbol)
                continue
            if latest >= target_session:
                current += 1
                continue
            if cached_rows < minimum_rows:
                bootstrap.append(symbol)
                continue
            cached = self._read_price_cache_unchecked(symbol)
            if cached is None:
                bootstrap.append(symbol)
                continue
            existing_frames[symbol] = cached
            pending.append(symbol)

        market_cfg = self.config.get("market_data", {})
        batch_size = int(market_cfg.get("batch_size", 50))
        overlap_days = int(market_cfg.get("incremental_overlap_days", 7))
        start_date = target_session - timedelta(days=max(2, overlap_days))
        end_date = self.now.date() + timedelta(days=1)
        refreshed = 0
        errors: list[dict[str, str]] = []

        for offset in range(0, len(pending), batch_size):
            batch = pending[offset : offset + batch_size]
            try:
                downloaded = self._download_batch_range(
                    batch,
                    start_date=start_date,
                    end_date=end_date,
                )
            except Exception as exc:
                errors.extend(
                    {"yahoo_symbol": symbol, "error": str(exc)} for symbol in batch
                )
                continue
            for symbol in batch:
                recent = downloaded.get(symbol)
                if recent is None or recent.empty:
                    errors.append(
                        {
                            "yahoo_symbol": symbol,
                            "error": "yfinance returned no recent daily prices",
                        }
                    )
                    continue
                existing = existing_frames.get(symbol)
                combined = recent if existing is None else pd.concat([existing, recent])
                combined = self._normalize_price_frame(combined)
                self._write_price_cache(symbol, combined)
                refreshed += 1

        bootstrap_period = str(market_cfg.get("incremental_bootstrap_period", "2y"))
        for offset in range(0, len(bootstrap), batch_size):
            batch = bootstrap[offset : offset + batch_size]
            try:
                downloaded = self._download_batch(batch, period=bootstrap_period)
            except Exception as exc:
                errors.extend(
                    {"yahoo_symbol": symbol, "error": str(exc)} for symbol in batch
                )
                continue
            for symbol in batch:
                frame = downloaded.get(symbol)
                if frame is None or frame.empty:
                    errors.append(
                        {
                            "yahoo_symbol": symbol,
                            "error": "yfinance returned no bootstrap price history",
                        }
                    )
                    continue
                self._write_price_cache(symbol, frame)
                refreshed += 1

        return SP500IncrementalRefreshResult(
            target_session=target_session,
            symbols_total=len(symbols),
            symbols_current=current,
            symbols_refreshed=refreshed,
            symbols_failed=len(errors),
            errors=errors,
        )

    def run(
        self,
        *,
        force: bool = False,
        refresh_constituents: bool = False,
        limit: int | None = None,
        period: str | None = None,
        batch_size: int | None = None,
    ) -> SP500ScoringResult:
        started_at = datetime.now(UTC)
        constituents = self.load_constituents(force=refresh_constituents)
        if limit is not None:
            constituents = constituents.head(limit).copy()

        market_cfg = self.config.get("market_data", {})
        selected_period = period or str(market_cfg.get("period", "5y"))
        selected_batch_size = batch_size or int(market_cfg.get("batch_size", 50))
        frames, download_errors, cache_hits = self.load_price_history(
            constituents["yahoo_symbol"].tolist(),
            force=force,
            period=selected_period,
            batch_size=selected_batch_size,
        )

        return self._build_scoring_result(
            constituents,
            frames,
            download_errors,
            started_at=started_at,
            metadata={
                "price_cache_hits": cache_hits,
                "history_period": selected_period,
            },
        )

    def run_cached(
        self,
        *,
        refresh_constituents: bool = False,
        limit: int | None = None,
    ) -> SP500ScoringResult:
        """Score the universe from local price caches without network downloads."""
        started_at = datetime.now(UTC)
        constituents = self.load_constituents(force=refresh_constituents)
        if limit is not None:
            constituents = constituents.head(limit).copy()
        frames: dict[str, pd.DataFrame] = {}
        errors: list[dict[str, Any]] = []
        for _, constituent in constituents.iterrows():
            yahoo_symbol = str(constituent["yahoo_symbol"])
            frame = self._read_price_cache_unchecked(yahoo_symbol)
            if frame is None or frame.empty:
                errors.append(self._error_row(constituent, "No cached price history available"))
                continue
            frames[yahoo_symbol] = frame

        return self._build_scoring_result(
            constituents,
            frames,
            errors,
            started_at=started_at,
            metadata={
                "price_cache_hits": len(frames),
                "history_period": "cached",
            },
        )

    def cached_price_history(self, yahoo_symbol: str) -> pd.DataFrame:
        """Return an isolated copy of a constituent's cached daily history."""
        frame = self._read_price_cache_unchecked(yahoo_symbol)
        return pd.DataFrame() if frame is None else frame.copy()

    def _build_scoring_result(
        self,
        constituents: pd.DataFrame,
        frames: dict[str, pd.DataFrame],
        initial_errors: list[dict[str, Any]],
        *,
        started_at: datetime,
        metadata: dict[str, Any],
    ) -> SP500ScoringResult:

        rows: list[dict[str, Any]] = []
        errors = list(initial_errors)
        minimum_rows = int(
            self.config.get("market_data", {}).get("min_history_rows", 220)
        )
        for index, constituent in constituents.reset_index(drop=True).iterrows():
            yahoo_symbol = str(constituent["yahoo_symbol"])
            frame = frames.get(yahoo_symbol)
            if frame is None or frame.empty:
                if not any(error["yahoo_symbol"] == yahoo_symbol for error in errors):
                    errors.append(self._error_row(constituent, "No price history available"))
                continue
            if len(frame) < minimum_rows:
                errors.append(
                    self._error_row(
                        constituent,
                        f"Insufficient history: {len(frame)} rows; minimum is {minimum_rows}",
                    )
                )
                continue
            try:
                rows.append(self.score_asset(constituent, frame, synthetic_id=-(index + 1)))
            except Exception as exc:  # One bad constituent must not abort the study.
                logger.exception("Scoring failed for %s", yahoo_symbol)
                errors.append(self._error_row(constituent, str(exc)))

        ranking = pd.DataFrame(rows)
        if not ranking.empty:
            ranking = ranking.sort_values(
                ["final_score", "technical_score"], ascending=[False, False]
            ).reset_index(drop=True)
            ranking.insert(0, "rank", range(1, len(ranking) + 1))

        finished_at = datetime.now(UTC)
        metadata = {
            **metadata,
            "started_at_utc": started_at.isoformat(),
            "finished_at_utc": finished_at.isoformat(),
            "duration_seconds": round((finished_at - started_at).total_seconds(), 2),
            "constituents_requested": len(constituents),
            "assets_scored": len(ranking),
            "assets_failed": len(errors),
            "price_provider": "yfinance",
            "portfolio_context": "neutral",
            "database_assets_modified": False,
        }
        return SP500ScoringResult(
            ranking=ranking,
            errors=pd.DataFrame(errors),
            metadata=metadata,
        )

    def load_constituents(self, *, force: bool = False) -> pd.DataFrame:
        cfg = self.config.get("constituents", {})
        cache_path = self._absolute_path(str(cfg.get("cache_file")))
        if cache_path.exists() and not force:
            return self._normalize_constituents(pd.read_csv(cache_path))

        url = str(cfg.get("source_url"))
        response = httpx.get(url, timeout=30, follow_redirects=True)
        response.raise_for_status()
        frame = self._normalize_constituents(pd.read_csv(StringIO(response.text)))
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        frame.to_csv(cache_path, index=False)
        return frame

    def load_price_history(
        self,
        yahoo_symbols: list[str],
        *,
        force: bool,
        period: str,
        batch_size: int,
    ) -> tuple[dict[str, pd.DataFrame], list[dict[str, Any]], int]:
        frames: dict[str, pd.DataFrame] = {}
        pending: list[str] = []
        cache_hits = 0
        for symbol in yahoo_symbols:
            cached = None if force else self._read_price_cache(symbol)
            if cached is not None:
                frames[symbol] = cached
                cache_hits += 1
            else:
                pending.append(symbol)

        errors: list[dict[str, Any]] = []
        for offset in range(0, len(pending), batch_size):
            symbols = pending[offset : offset + batch_size]
            logger.info(
                "Downloading S&P 500 price batch %s-%s of %s",
                offset + 1,
                min(offset + len(symbols), len(pending)),
                len(pending),
            )
            try:
                downloaded = self._download_batch(symbols, period=period)
            except Exception as exc:
                logger.warning("Price batch failed: %s", exc)
                errors.extend(
                    {
                        "symbol": self._display_symbol(symbol),
                        "yahoo_symbol": symbol,
                        "error": str(exc),
                    }
                    for symbol in symbols
                )
                continue
            for symbol in symbols:
                frame = downloaded.get(symbol)
                if frame is None or frame.empty:
                    errors.append(
                        {
                            "symbol": self._display_symbol(symbol),
                            "yahoo_symbol": symbol,
                            "error": "yfinance returned no daily price history",
                        }
                    )
                    continue
                frames[symbol] = frame
                self._write_price_cache(symbol, frame)
        return frames, errors, cache_hits

    def _download_batch(self, symbols: list[str], *, period: str) -> dict[str, pd.DataFrame]:
        import yfinance as yf

        cfg = self.config.get("market_data", {})
        attempts = max(1, int(cfg.get("retries", 2)))
        last_error: Exception | None = None
        for attempt in range(attempts):
            try:
                raw = yf.download(
                    tickers=symbols,
                    period=period,
                    interval=str(cfg.get("interval", "1d")),
                    group_by="ticker",
                    auto_adjust=False,
                    threads=True,
                    progress=False,
                    timeout=30,
                )
                return {
                    symbol: frame
                    for symbol in symbols
                    if not (frame := self._extract_symbol_frame(raw, symbol, len(symbols))).empty
                }
            except Exception as exc:
                last_error = exc
                if attempt + 1 < attempts:
                    time.sleep(float(cfg.get("retry_backoff_seconds", 3)) * (attempt + 1))
        raise RuntimeError(f"yfinance batch failed after {attempts} attempts: {last_error}")

    def _download_batch_range(
        self,
        symbols: list[str],
        *,
        start_date: date,
        end_date: date,
    ) -> dict[str, pd.DataFrame]:
        import yfinance as yf

        cfg = self.config.get("market_data", {})
        attempts = max(1, int(cfg.get("retries", 2)))
        last_error: Exception | None = None
        for attempt in range(attempts):
            try:
                raw = yf.download(
                    tickers=symbols,
                    start=start_date.isoformat(),
                    end=end_date.isoformat(),
                    interval=str(cfg.get("interval", "1d")),
                    group_by="ticker",
                    auto_adjust=False,
                    threads=True,
                    progress=False,
                    timeout=30,
                )
                return {
                    symbol: frame
                    for symbol in symbols
                    if not (
                        frame := self._extract_symbol_frame(raw, symbol, len(symbols))
                    ).empty
                }
            except Exception as exc:
                last_error = exc
                if attempt + 1 < attempts:
                    time.sleep(float(cfg.get("retry_backoff_seconds", 3)) * (attempt + 1))
        raise RuntimeError(
            f"yfinance incremental batch failed after {attempts} attempts: {last_error}"
        )

    def score_asset(
        self,
        constituent: pd.Series,
        frame: pd.DataFrame,
        *,
        synthetic_id: int,
    ) -> dict[str, Any]:
        min_rows = int(self.config.get("market_data", {}).get("min_history_rows", 220))
        if len(frame) < min_rows:
            raise ValueError(f"Insufficient history: {len(frame)} rows; minimum is {min_rows}")

        enriched = self.technical.compute_indicators(frame)
        latest = enriched.iloc[-1]
        support = self.support.detect_support_zone(enriched)
        trend_score, trend_rationale = self.technical.trend_structure_score(latest)
        technical_score, technical_rationale = self.scoring.compute_technical_score(
            latest,
            distance_to_support_pct=support.distance_to_support_pct,
            trend_score=trend_score,
        )
        asset = AssetORM(
            id=synthetic_id,
            symbol=str(constituent["symbol"]),
            name=str(constituent["name"]),
            asset_type="stock",
            sector=str(constituent["sector"]),
            region="US",
            enabled=False,
            supports_fundamentals=False,
        )
        portfolio_fit, _ = self.rebalance.portfolio_fit_score(asset, self._neutral_exposure())
        risk = self.risk.assess_risk(enriched, asset_type="stock")
        final_details = self.scoring.compute_final_score_details(
            technical_score=technical_score,
            risk_score=risk.risk_score,
            portfolio_fit_score=portfolio_fit,
        )
        snapshot = TechnicalSnapshotModel(
            asset_id=synthetic_id,
            date=pd.Timestamp(latest["date"]).date(),
            rsi14=self._optional_float(latest.get("rsi14")),
            sma50=self._optional_float(latest.get("sma50")),
            sma200=self._optional_float(latest.get("sma200")),
            ema20=self._optional_float(latest.get("ema20")),
            atr14=self._optional_float(latest.get("atr14")),
            week_52_low=self._optional_float(latest.get("week_52_low")),
            week_52_high=self._optional_float(latest.get("week_52_high")),
            week_52_position=self._optional_float(technical_rationale.get("week_52_position")),
            distance_52w_high_pct=self._optional_float(latest.get("distance_52w_high_pct")),
            distance_52w_low_pct=self._optional_float(latest.get("distance_52w_low_pct")),
            support_low=support.support_zone_low,
            support_high=support.support_zone_high,
            distance_to_support_pct=support.distance_to_support_pct,
            technical_score=technical_score,
            rationale={**trend_rationale, **technical_rationale},
        )
        recommendation = self.recommendation.build_recommendation(
            asset_id=synthetic_id,
            as_of_date=snapshot.date,
            technical_snapshot=snapshot,
            risk=risk,
            final_score=float(final_details["final"]),
            portfolio_fit_score=portfolio_fit,
            score_breakdown=final_details,
        )

        regime_values: dict[str, Any] = {
            "bull_probability": None,
            "bear_probability": None,
            "bubble_probability": None,
            "dominant_regime": None,
        }
        if self.config.get("scoring", {}).get("include_market_regime", True):
            regime = self.regime.compute_regime(enriched)
            regime_values = {
                "bull_probability": regime.bull_probability,
                "bear_probability": regime.bear_probability,
                "bubble_probability": regime.bubble_probability,
                "dominant_regime": regime.dominant_regime,
            }

        return {
            "symbol": str(constituent["symbol"]),
            "yahoo_symbol": str(constituent["yahoo_symbol"]),
            "company": str(constituent["name"]),
            "sector": str(constituent["sector"]),
            "sub_industry": str(constituent.get("sub_industry", "")),
            "price_date": snapshot.date,
            "last_price": round(float(latest["close"]), 4),
            "rsi14": self._rounded(latest.get("rsi14")),
            "sma50": self._rounded(latest.get("sma50")),
            "sma200": self._rounded(latest.get("sma200")),
            "distance_52w_high_pct": self._rounded(latest.get("distance_52w_high_pct")),
            "support_low": support.support_zone_low,
            "support_high": support.support_zone_high,
            "distance_to_support_pct": support.distance_to_support_pct,
            "technical_score": round(float(technical_score), 2),
            "risk_score": round(float(risk.risk_score), 2),
            "risk_level": risk.risk_level.value,
            "portfolio_fit_score": round(float(portfolio_fit), 2),
            "final_score": round(float(final_details["final"]), 2),
            "recommendation": recommendation.recommendation.value,
            "suggested_weight_add_pct": recommendation.suggested_weight_add,
            "invalidation": recommendation.invalidation,
            **regime_values,
            "data_provider": "yfinance",
            "history_rows": len(frame),
            "history_start": pd.Timestamp(frame["date"].min()).date(),
        }

    def export_excel(
        self,
        result: SP500ScoringResult,
        *,
        output_path: Path | None = None,
    ) -> Path:
        if output_path is None:
            cfg = self.config.get("output", {})
            output_dir = self._absolute_path(str(cfg.get("directory", "outputs/screeners")))
            prefix = str(cfg.get("filename_prefix", "sp500_scoring"))
            output_path = output_dir / f"{prefix}_{date.today().isoformat()}.xlsx"
        output_path.parent.mkdir(parents=True, exist_ok=True)

        metadata = pd.DataFrame(
            [
                {"key": key, "value": self._excel_value(value)}
                for key, value in result.metadata.items()
            ]
        )
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            result.ranking.to_excel(writer, sheet_name="Ranking", index=False)
            result.errors.to_excel(writer, sheet_name="Errors", index=False)
            metadata.to_excel(writer, sheet_name="Metadata", index=False)
            self._style_workbook(writer.book)
        return output_path

    def _read_price_cache(self, symbol: str) -> pd.DataFrame | None:
        path = self._price_cache_path(symbol)
        if not path.exists():
            return None
        max_age = float(self.config.get("market_data", {}).get("cache_max_age_hours", 18))
        age_hours = (time.time() - path.stat().st_mtime) / 3600
        if age_hours > max_age:
            return None
        return self._read_price_cache_unchecked(symbol)

    def _read_price_cache_unchecked(self, symbol: str) -> pd.DataFrame | None:
        path = self._price_cache_path(symbol)
        if not path.exists():
            return None
        try:
            return self._normalize_price_frame(pd.read_csv(path))
        except Exception:
            logger.warning("Ignoring invalid price cache for %s", symbol, exc_info=True)
            return None

    def _cached_price_status(self, symbol: str) -> tuple[date | None, int]:
        path = self._price_cache_path(symbol)
        if not path.exists():
            return None, 0
        try:
            values = pd.to_datetime(
                pd.read_csv(path, usecols=["date"])["date"], errors="coerce"
            ).dropna()
            if values.empty:
                return None, 0
            return values.max().date(), len(values)
        except Exception:
            logger.warning("Cannot inspect price cache for %s", symbol, exc_info=True)
            return None, 0

    def _expected_market_session(self) -> date:
        today = self.now.date()
        cutoff_hour = int(
            self.config.get("market_data", {}).get("market_close_cutoff_utc_hour", 21)
        )
        if self.now.hour >= cutoff_hour and pd.Timestamp(today).dayofweek < 5:
            return today
        return (pd.Timestamp(today) - pd.offsets.BDay(1)).date()

    def _write_price_cache(self, symbol: str, frame: pd.DataFrame) -> None:
        path = self._price_cache_path(symbol)
        path.parent.mkdir(parents=True, exist_ok=True)
        frame.to_csv(path, index=False)

    def _price_cache_path(self, symbol: str) -> Path:
        directory = self._absolute_path(
            str(self.config.get("market_data", {}).get("cache_dir", "cache/sp500_scoring/prices"))
        )
        return directory / f"{symbol.replace('/', '_')}.csv"

    @classmethod
    def _extract_symbol_frame(
        cls, raw: pd.DataFrame, symbol: str, symbol_count: int
    ) -> pd.DataFrame:
        if raw.empty:
            return pd.DataFrame()
        if isinstance(raw.columns, pd.MultiIndex):
            if symbol in raw.columns.get_level_values(0):
                selected = raw[symbol].copy()
            elif symbol in raw.columns.get_level_values(1):
                selected = raw.xs(symbol, axis=1, level=1).copy()
            else:
                return pd.DataFrame()
        elif symbol_count == 1:
            selected = raw.copy()
        else:
            return pd.DataFrame()
        selected = selected.reset_index()
        return cls._normalize_price_frame(selected)

    @staticmethod
    def _normalize_price_frame(frame: pd.DataFrame) -> pd.DataFrame:
        result = frame.copy()
        result.columns = [
            str(column).strip().lower().replace(" ", "_") for column in result.columns
        ]
        if "datetime" in result.columns and "date" not in result.columns:
            result = result.rename(columns={"datetime": "date"})
        required = ["date", "open", "high", "low", "close"]
        if any(column not in result.columns for column in required):
            return pd.DataFrame()
        result["date"] = pd.to_datetime(
            result["date"], utc=True, errors="coerce"
        ).dt.tz_localize(None)
        for column in ["open", "high", "low", "close", "volume"]:
            if column not in result.columns:
                result[column] = 0.0
            result[column] = pd.to_numeric(result[column], errors="coerce")
        result = result.dropna(subset=["date", "close"])
        result = result[result["close"] > 0]
        result["volume"] = result["volume"].fillna(0.0)
        return result[["date", "open", "high", "low", "close", "volume"]].sort_values(
            "date"
        ).drop_duplicates("date", keep="last").reset_index(drop=True)

    @classmethod
    def _normalize_constituents(cls, frame: pd.DataFrame) -> pd.DataFrame:
        aliases = {
            "Symbol": "symbol",
            "Security": "name",
            "GICS Sector": "sector",
            "GICS Sub-Industry": "sub_industry",
        }
        result = frame.rename(columns=aliases).copy()
        required = {"symbol", "name", "sector"}
        missing = required.difference(result.columns)
        if missing:
            raise ValueError(f"Constituent source is missing columns: {sorted(missing)}")
        if "sub_industry" not in result.columns:
            result["sub_industry"] = ""
        result["symbol"] = result["symbol"].astype(str).str.strip().str.upper()
        result["yahoo_symbol"] = result["symbol"].map(cls.to_yahoo_symbol)
        return result[["symbol", "yahoo_symbol", "name", "sector", "sub_industry"]].drop_duplicates(
            "symbol"
        )

    @staticmethod
    def to_yahoo_symbol(symbol: str) -> str:
        return symbol.strip().upper().replace(".", "-")

    @staticmethod
    def _display_symbol(yahoo_symbol: str) -> str:
        return yahoo_symbol.replace("-", ".") if yahoo_symbol in {"BRK-B", "BF-B"} else yahoo_symbol

    @staticmethod
    def _neutral_exposure() -> PortfolioExposureModel:
        return PortfolioExposureModel(
            total_invested_weight=0.0,
            by_asset={},
            by_sector={},
            by_asset_type={},
        )

    @staticmethod
    def _optional_float(value: Any) -> float | None:
        if value is None or pd.isna(value):
            return None
        return float(value)

    @classmethod
    def _rounded(cls, value: Any) -> float | None:
        number = cls._optional_float(value)
        return round(number, 4) if number is not None else None

    @staticmethod
    def _error_row(constituent: pd.Series, error: str) -> dict[str, Any]:
        return {
            "symbol": str(constituent["symbol"]),
            "yahoo_symbol": str(constituent["yahoo_symbol"]),
            "company": str(constituent["name"]),
            "error": error,
        }

    def _absolute_path(self, configured_path: str) -> Path:
        path = Path(configured_path)
        return path if path.is_absolute() else self.root_dir / path

    @staticmethod
    def _excel_value(value: Any) -> Any:
        return json.dumps(value, ensure_ascii=True) if isinstance(value, dict | list) else value

    @staticmethod
    def _style_workbook(workbook: Any) -> None:
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Font, PatternFill

        header_fill = PatternFill("solid", fgColor="17324D")
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
            for column in worksheet.columns:
                values = [
                    len(str(cell.value)) if cell.value is not None else 0
                    for cell in column[:100]
                ]
                width = min(max(values + [10]) + 2, 45)
                worksheet.column_dimensions[column[0].column_letter].width = width
        ranking = workbook["Ranking"]
        headers = {cell.value: cell.column_letter for cell in ranking[1]}
        for score_column in ["technical_score", "final_score"]:
            letter = headers.get(score_column)
            if letter and ranking.max_row > 1:
                ranking.conditional_formatting.add(
                    f"{letter}2:{letter}{ranking.max_row}",
                    ColorScaleRule(
                        start_type="min",
                        start_color="F8696B",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FFEB84",
                        end_type="max",
                        end_color="63BE7B",
                    ),
                )
        risk_letter = headers.get("risk_score")
        if risk_letter and ranking.max_row > 1:
            ranking.conditional_formatting.add(
                f"{risk_letter}2:{risk_letter}{ranking.max_row}",
                ColorScaleRule(
                    start_type="min",
                    start_color="63BE7B",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color="F8696B",
                ),
            )
